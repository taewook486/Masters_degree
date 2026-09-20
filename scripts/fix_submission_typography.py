"""재빌드한 제출본 docx의 조판 결함을 정정한다.

배경: `build_thesis_docx.py` + `restore_table_formatting.py`를 거친 docx에는
두 가지 조판 문제가 남는다.

1. 열이 많은 표(Table 4.3a는 12열)에서 머리글이 줄바꿈된다. 기본 셀 여백이
   좌우 108twips씩이라 12열이면 2,592twips(본문 가용폭 7,482의 35%)가 여백으로
   나가고, 남은 자리에 굵은 머리글이 들어가지 못한다. 2026-09-18 지도교수
   지적의 `warmup` → `war`/`mup` 갈라짐이 이 현상이다.
2. 국문 학교 양식의 본문 구역이 `w:pgNumType w:fmt="decimalFullWidth"`라
   쪽번호와 목차 쪽수가 전각(４９)으로 찍힌다. 목차에서 단어와 쪽수 사이가
   어긋나 보이는 원인이다. 영문 양식은 원래 일반 숫자라 손대지 않는다.

정정 방식: 셀 여백을 좁히고, 남은 폭을 열별 실제 필요폭에 비례해 재분배한다.
표 전체 폭은 건드리지 않으므로 양식의 본문 영역을 벗어나지 않는다.

사용:
    python3 scripts/fix_submission_typography.py --in final_ko.docx --lang ko
"""

from __future__ import annotations

import argparse
import re
import unicodedata
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.shared import Twips

# 9pt·장평 97% 기준 글자당 폭(twips). 현 제출본의 실측 열 폭에서 역산했다.
# 굵은 머리글이 일반 본문보다 넓어 두 값을 나눠 쓴다.
TW_ASCII_BOLD = 62
TW_ASCII = 53
TW_CJK = 175

CELL_MARGIN_TW = 28  # 좌우 각각. Word 기본값은 108이다.
SAFETY_TW = 24  # 글자폭 추정 오차 흡수분

# 여백을 좁힐 표의 최소 열 수. 열이 적은 표는 기본 여백이 보기 좋다.
WIDE_TABLE_MIN_COLS = 6


def _char_width(ch: str, bold: bool) -> int:
    if unicodedata.east_asian_width(ch) in ("W", "F"):
        return TW_CJK
    return TW_ASCII_BOLD if bold else TW_ASCII


def _text_width(text: str, bold: bool) -> int:
    return sum(_char_width(c, bold) for c in text)


def _needed_width(table, col_idx: int) -> int:
    """열이 줄바꿈 없이 담아야 하는 최대 폭."""
    widest = 0
    for row_idx, row in enumerate(table.rows):
        cell = row.cells[col_idx]
        bold = row_idx == 0  # 머리행은 굵게 조판된다
        for para in cell.paragraphs:
            widest = max(widest, _text_width(para.text.strip(), bold))
    return widest + SAFETY_TW


def _set_cell_margins(table, margin_tw: int) -> None:
    tbl_pr = table._tbl.tblPr
    existing = tbl_pr.find(qn("w:tblCellMar"))
    if existing is not None:
        tbl_pr.remove(existing)
    mar = tbl_pr.makeelement(qn("w:tblCellMar"), {})
    for side, value in (("left", margin_tw), ("right", margin_tw)):
        node = mar.makeelement(qn(f"w:{side}"), {})
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tbl_pr.append(mar)


def _redistribute(table, margin_tw: int) -> tuple[int, int]:
    """열 폭을 필요폭에 비례해 재분배한다. 표 전체 폭은 유지한다."""
    total = sum(c.width.twips for c in table.columns)
    n = len(table.columns)
    budget = total - n * 2 * margin_tw  # 글자가 실제로 쓸 수 있는 합

    needs = [_needed_width(table, i) for i in range(n)]
    need_sum = sum(needs)

    if need_sum <= budget:
        # 여유가 있으면 필요폭을 보장하고 남는 폭만 비례 배분한다.
        slack = budget - need_sum
        alloc = [nd + round(slack * nd / need_sum) for nd in needs]
    else:
        # 모자라면 전체를 필요폭 비례로 눌러 담는다(줄바꿈은 생기되 고르게).
        alloc = [round(budget * nd / need_sum) for nd in needs]

    # 반올림 오차를 가장 넓은 열에서 흡수해 합계를 정확히 맞춘다.
    drift = budget - sum(alloc)
    alloc[alloc.index(max(alloc))] += drift

    for col, text_w in zip(table.columns, alloc):
        col.width = Twips(text_w + 2 * margin_tw)
    return total, sum(c.width.twips for c in table.columns)


def _fix_page_numbers(doc) -> int:
    """전각 쪽번호를 일반 숫자로 바꾼다. 바꾼 구역 수를 돌려준다."""
    changed = 0
    for section in doc.sections:
        pg = section._sectPr.find(qn("w:pgNumType"))
        if pg is None:
            continue
        if pg.get(qn("w:fmt")) == "decimalFullWidth":
            pg.set(qn("w:fmt"), "decimal")
            changed += 1
    return changed


def _caption_before(doc, table_index: int) -> str:
    from docx.text.paragraph import Paragraph

    last = ""
    seen = 0
    for child in doc.element.body:
        if child.tag.endswith("}p"):
            text = Paragraph(child, doc).text.strip()
            if text:
                last = text
        elif child.tag.endswith("}tbl"):
            if seen == table_index:
                match = re.match(r"(Table [0-9A-Za-z.]+)", last)
                return match.group(1) if match else last[:24]
            seen += 1
    return ""


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--in", dest="inp", required=True)
    ap.add_argument("--out", default=None, help="생략하면 입력 파일을 덮어쓴다")
    ap.add_argument("--lang", choices=["ko", "en"], required=True)
    args = ap.parse_args()

    path = Path(args.inp)
    doc = docx.Document(str(path))

    touched = 0
    for idx, table in enumerate(doc.tables):
        if len(table.columns) < WIDE_TABLE_MIN_COLS:
            continue
        caption = _caption_before(doc, idx)
        before, after = _redistribute(table, CELL_MARGIN_TW)
        _set_cell_margins(table, CELL_MARGIN_TW)
        touched += 1
        label = caption or f"표{idx}"
        cols = len(table.columns)
        print(f"  [{label}] {cols}열 폭 {before} → {after}tw")

    if args.lang == "ko":
        n = _fix_page_numbers(doc)
        print(f"  쪽번호 전각 → 일반: {n}개 구역")
    else:
        print("  쪽번호: 영문 양식은 원래 일반 숫자 — 변경 없음")

    out = Path(args.out) if args.out else path
    doc.save(str(out))
    print(f"저장: {out}  (표 {touched}개 조정)")


if __name__ == "__main__":
    main()
